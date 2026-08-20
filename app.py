from flask import Flask, render_template, request, jsonify, send_file
import os
import io
import webbrowser
import threading
import time
import logging
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db
from recordatorio import gestor_recordatorio

# Desactivar logs molestos de peticiones HTTP en la consola de terminal
log = logging.getLogger('werkzeug')
log.setLevel(logging.ERROR)

import sys
if getattr(sys, 'frozen', False):
    template_folder = os.path.join(sys._MEIPASS, 'templates')
    static_folder = os.path.join(sys._MEIPASS, 'static')
    app = Flask(__name__, template_folder=template_folder, static_folder=static_folder)
    app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024 # 50 MB limit
else:
    app = Flask(__name__)
    app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024 # 50 MB limit
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

# Inicializar BD y Recordatorio al arrancar
with app.app_context():
    db.init_db()
    # Iniciar recordatorio dinámico con configuración persistente
    gestor_recordatorio.iniciar()

@app.context_processor
def inject_version():
    return dict(cache_v=int(time.time()))

@app.after_request
def add_cors_headers(response):
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type,Authorization'
    response.headers['Access-Control-Allow-Methods'] = 'GET,PUT,POST,DELETE,OPTIONS'
    return response

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/recordatorio', methods=['GET'])
def get_recordatorio():
    return jsonify({"status": "success", "data": gestor_recordatorio.estado()})

@app.route('/api/recordatorio/audio', methods=['POST'])
def upload_audio():
    if 'audio' not in request.files:
        return jsonify({"status": "error", "message": "No se envió ningún archivo de audio."}), 400
    
    file = request.files['audio']
    if file.filename == '':
        return jsonify({"status": "error", "message": "Archivo no seleccionado."}), 400
        
    alarma_id = request.form.get('alarma_id', 'alerta')
    if not alarma_id:
        alarma_id = 'alerta'
        
    if not (file.filename.lower().endswith('.wav') or file.filename.lower().endswith('.mp3')):
        return jsonify({"status": "error", "message": "El archivo debe ser formato .wav o .mp3."}), 400
        
    try:
        import os
        import ctypes
        # Cerrar cualquier reproduccion activa para evitar error de archivo bloqueado
        ctypes.windll.winmm.mciSendStringW('close mi_audio', None, 0, None)
        
        base_dir = os.path.dirname(os.path.abspath(__file__))
        ext = '.mp3' if file.filename.lower().endswith('.mp3') else '.wav'
        # Eliminar el archivo con la extension opuesta si existe para evitar conflictos
        old_ext = '.wav' if ext == '.mp3' else '.mp3'
        old_path = os.path.join(base_dir, f"{alarma_id}{old_ext}")
        if os.path.exists(old_path):
            os.remove(old_path)
            
        filename = f"{alarma_id}{ext}"
        save_path = os.path.join(base_dir, filename)
        # Si el mismo archivo existe, intentar removerlo
        if os.path.exists(save_path):
            os.remove(save_path)
        file.save(save_path)
        return jsonify({"status": "success", "message": f"Audio subido correctamente para la alarma."})
    except Exception as e:
        return jsonify({"status": "error", "message": f"Error al guardar audio: {str(e)}"}), 500


@app.route('/api/recordatorio', methods=['POST'])
def controlar_recordatorio():
    data = request.json
    accion = data.get('accion')
    
    if accion == 'test':
        alarma_id = data.get('alarma_id')
        if gestor_recordatorio.probar_alarma(alarma_id):
            return jsonify({"status": "success", "message": "Notificación de prueba enviada."})
        return jsonify({"status": "error", "message": "No se pudo enviar la alarma o no existe el ID."})
    elif accion == 'guardar_config':
        nueva_config = data.get('config')
        if not nueva_config:
            return jsonify({"status": "error", "message": "No se recibió configuración."}), 400
        gestor_recordatorio.guardar_config(nueva_config)
        return jsonify({"status": "success", "data": gestor_recordatorio.estado(), "message": "¡Configuración de alarmas guardada y aplicada!"})
    elif accion == 'toggle':
        if gestor_recordatorio.activo:
            gestor_recordatorio.detener()
            msg = "Recordatorios automáticos desactivados."
        else:
            gestor_recordatorio.iniciar()
            msg = "Recordatorios automáticos activados."
        return jsonify({"status": "success", "data": gestor_recordatorio.estado(), "message": msg})
    
    return jsonify({"status": "error", "message": "Acción no reconocida."}), 400

@app.route('/api/sugerencias', methods=['GET'])
def get_sugerencias():
    sugerencias = db.get_sugerencias()
    return jsonify({"status": "success", "data": sugerencias})

@app.route('/api/feriados', methods=['GET'])
def get_feriados():
    anio = request.args.get('year', type=int)
    feriados = db.get_feriados_chile(anio)
    return jsonify({"status": "success", "data": feriados})

# ==========================================
# ENDPOINTS API AGENDA Y PLANNER DE TAREAS
# ==========================================

@app.route('/api/agenda', methods=['GET'])
def get_agenda_tareas():
    fecha = request.args.get('fecha')
    estado = request.args.get('estado')
    solo_proximas = request.args.get('solo_proximas', '').lower() == 'true'
    tareas = db.get_all_agenda_tareas(fecha, estado, solo_proximas)
    return jsonify({"status": "success", "data": tareas})

@app.route('/api/agenda', methods=['POST'])
def add_agenda_tarea():
    data = request.json
    try:
        titulo = data.get('titulo', '').strip()
        if not titulo:
            return jsonify({"status": "error", "message": "El título de la tarea es obligatorio."}), 400
            
        descripcion = data.get('descripcion', '').strip()
        fecha = data.get('fecha', datetime.now().strftime('%Y-%m-%d'))
        hora = data.get('hora', '09:00').strip()
        minutos_recordatorio = int(data.get('minutos_recordatorio', 10))
        lugar = data.get('lugar', '').strip()
        horas_estimadas = float(data.get('horas_estimadas', 1.0))
        estado = data.get('estado', 'pendiente').strip()
        
        nueva = db.add_agenda_tarea(titulo, descripcion, fecha, hora, minutos_recordatorio, lugar, horas_estimadas, estado)
        return jsonify({"status": "success", "data": nueva, "message": "Tarea agendada con éxito."}), 201
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/agenda/<int:tarea_id>', methods=['PUT'])
def update_agenda_tarea(tarea_id):
    data = request.json
    try:
        titulo = data.get('titulo', '').strip()
        if not titulo:
            return jsonify({"status": "error", "message": "El título de la tarea es obligatorio."}), 400
            
        descripcion = data.get('descripcion', '').strip()
        fecha = data.get('fecha', datetime.now().strftime('%Y-%m-%d'))
        hora = data.get('hora', '09:00').strip()
        minutos_recordatorio = int(data.get('minutos_recordatorio', 10))
        lugar = data.get('lugar', '').strip()
        horas_estimadas = float(data.get('horas_estimadas', 1.0))
        estado = data.get('estado', 'pendiente').strip()
        
        actualizada = db.update_agenda_tarea(tarea_id, titulo, descripcion, fecha, hora, minutos_recordatorio, lugar, horas_estimadas, estado)
        return jsonify({"status": "success", "data": actualizada, "message": "Tarea actualizada correctamente."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/agenda/<int:tarea_id>/estado', methods=['PATCH'])
def update_estado_agenda_tarea(tarea_id):
    data = request.json
    try:
        nuevo_estado = data.get('estado', 'pendiente').strip()
        actualizada = db.update_estado_agenda_tarea(tarea_id, nuevo_estado)
        return jsonify({"status": "success", "data": actualizada, "message": f"Estado cambiado a {nuevo_estado}."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/agenda/<int:tarea_id>', methods=['DELETE'])
def delete_agenda_tarea(tarea_id):
    try:
        db.delete_agenda_tarea(tarea_id)
        return jsonify({"status": "success", "message": "Tarea eliminada."})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/agenda/<int:tarea_id>/convertir_a_registro', methods=['POST'])
def convertir_tarea_a_registro(tarea_id):
    data = request.json or {}
    try:
        tarea = db.get_agenda_tarea(tarea_id)
        if not tarea:
            return jsonify({"status": "error", "message": "Tarea no encontrada."}), 404
            
        fecha = data.get('fecha') or tarea['fecha']
        lugar = data.get('lugar') or tarea['lugar'] or 'Oficina'
        horas = float(data.get('horas') or tarea['horas_estimadas'] or 1.0)
        actividad = data.get('actividad') or tarea['titulo']
        comentario = data.get('comentario') or tarea['descripcion'] or 'Tarea agendada completada'
        
        nuevo_reg = db.add_registro(fecha, lugar, horas, actividad, comentario)
        db.update_estado_agenda_tarea(tarea_id, 'completada')
        
        return jsonify({"status": "success", "data": nuevo_reg, "message": "¡Tarea completada y registrada en tus horas trabajadas!"}), 201
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/registros', methods=['GET'])
def get_registros():
    fecha_inicio = request.args.get('fecha_inicio')
    fecha_fin = request.args.get('fecha_fin')
    lugar = request.args.get('lugar')
    actividad = request.args.get('actividad')
    search = request.args.get('search')
    
    registros = db.get_all_registros(fecha_inicio, fecha_fin, lugar, actividad, search)
    return jsonify({"status": "success", "data": registros})

@app.route('/api/registros', methods=['POST'])
def add_registro():
    data = request.json
    try:
        fecha = data.get('fecha', datetime.now().strftime('%Y-%m-%d'))
        lugar = data.get('lugar', '').strip()
        horas = float(data.get('horas', 0))
        actividad = data.get('actividad', '').strip()
        comentario = data.get('comentario', '').strip()
        
        if not lugar or not actividad or horas <= 0:
            return jsonify({"status": "error", "message": "Por favor completa el lugar, la actividad y las horas (mayor a 0)."}), 400
            
        nuevo = db.add_registro(fecha, lugar, horas, actividad, comentario)
        return jsonify({"status": "success", "data": nuevo, "message": "Registro guardado exitosamente."}), 201
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/registros/<int:registro_id>', methods=['PUT'])
def update_registro(registro_id):
    data = request.json
    try:
        fecha = data.get('fecha')
        lugar = data.get('lugar', '').strip()
        horas = float(data.get('horas', 0))
        actividad = data.get('actividad', '').strip()
        comentario = data.get('comentario', '').strip()
        
        if not lugar or not actividad or horas <= 0:
            return jsonify({"status": "error", "message": "Datos inválidos."}), 400
            
        actualizado = db.update_registro(registro_id, fecha, lugar, horas, actividad, comentario)
        if actualizado:
            return jsonify({"status": "success", "data": actualizado, "message": "Registro actualizado exitosamente."})
        else:
            return jsonify({"status": "error", "message": "Registro no encontrado."}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 400

@app.route('/api/registros/<int:registro_id>', methods=['DELETE'])
def delete_registro(registro_id):
    exito = db.delete_registro(registro_id)
    if exito:
        return jsonify({"status": "success", "message": "Registro eliminado."})
    else:
        return jsonify({"status": "error", "message": "No se pudo eliminar el registro."}), 404

@app.route('/api/dashboard', methods=['GET'])
def get_dashboard():
    stats = db.get_dashboard_stats()
    return jsonify({"status": "success", "data": stats})

@app.route('/api/export/excel', methods=['GET'])
def export_excel():
    fecha_inicio = request.args.get('fecha_inicio')
    fecha_fin = request.args.get('fecha_fin')
    lugar = request.args.get('lugar')
    actividad = request.args.get('actividad')
    search = request.args.get('search')
    
    registros = db.get_all_registros(fecha_inicio, fecha_fin, lugar, actividad, search)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro de Horas"
    
    title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="3B82F6", end_color="1D4ED8", fill_type="solid")
    
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    
    data_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = "REPORTE DE CONTROL DE HORAS DE TRABAJO"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35
    
    ws["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    
    headers = ["ID", "Fecha", "Lugar / Cliente / Proyecto", "Actividad", "Horas", "Comentario Complementario"]
    for col_num, header_text in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_num in [1, 2, 5] else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 24
    
    row_idx = 5
    feriados_list = db.get_feriados_chile()
    feriados_map = {f["fecha"]: f for f in feriados_list}
    holiday_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    holiday_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    for reg in registros:
        ws.cell(row=row_idx, column=1, value=reg["id"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=reg["fecha"]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=reg["lugar"])
        ws.cell(row=row_idx, column=4, value=reg["actividad"])
        
        horas_cell = ws.cell(row=row_idx, column=5, value=reg["horas"])
        horas_cell.number_format = '#,##0.00'
        horas_cell.alignment = Alignment(horizontal="center")
        
        comentario_texto = reg["comentario"] or ""
        f_info = feriados_map.get(reg["fecha"])
        if f_info:
            prefix = f"[🇨🇱 Feriado: {f_info['titulo']}]"
            comentario_texto = f"{prefix} {comentario_texto}".strip() if comentario_texto else prefix
            
        ws.cell(row=row_idx, column=6, value=comentario_texto)
        
        for col_num in range(1, 7):
            c = ws.cell(row=row_idx, column=col_num)
            c.font = data_font
            c.border = thin_border
            if row_idx % 2 == 0:
                c.fill = alt_fill
            if col_num == 2 and f_info:
                c.fill = holiday_fill
                c.font = holiday_font
                
        ws.row_dimensions[row_idx].height = 20
        row_idx += 1
        
    if len(registros) > 0:
        ws.merge_cells(f"A{row_idx}:D{row_idx}")
        total_label_cell = ws.cell(row=row_idx, column=1, value="TOTAL HORAS REGISTRADAS")
        total_label_cell.font = bold_font
        total_label_cell.alignment = Alignment(horizontal="right", vertical="center")
        
        total_sum_cell = ws.cell(row=row_idx, column=5, value=f"=SUM(E5:E{row_idx-1})")
        total_sum_cell.font = bold_font
        total_sum_cell.number_format = '#,##0.00'
        total_sum_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col_num in range(1, 7):
            c = ws.cell(row=row_idx, column=col_num)
            c.fill = total_fill
            c.border = Border(top=Side(style='medium', color='64748B'), bottom=Side(style='double', color='1E293B'))
        ws.row_dimensions[row_idx].height = 25
        
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1 and col_letter == 'A':
                continue
            try:
                if len(str(cell.value or '')) > max_len:
                    max_len = len(str(cell.value or ''))
            except:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 45
    
    # ---------------------------------------------------------
    # HOJA 2: RESUMEN AGRUPADO POR ACTIVIDAD (SIN REPETIR ACTIVIDADES)
    # ---------------------------------------------------------
    ws2 = wb.create_sheet(title="Horas por Actividad")
    
    title_fill2 = PatternFill(start_color="10B981", end_color="059669", fill_type="solid")
    ws2.merge_cells("A1:E1")
    title_cell2 = ws2["A1"]
    title_cell2.value = "RESUMEN AGRUPADO POR ACTIVIDAD (SIN REPETIR)"
    title_cell2.font = title_font
    title_cell2.fill = title_fill2
    title_cell2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 35
    
    ws2["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Vista consolidada por Actividad"
    ws2["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    
    headers2 = ["Actividad", "Clientes / Proyectos Involucrados", "Días / Sesiones", "Horas Totales", "% del Total"]
    for col_num, header_text in enumerate(headers2, 1):
        cell = ws2.cell(row=4, column=col_num)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_num in [3, 4, 5] else "left", vertical="center")
        cell.border = thin_border
    ws2.row_dimensions[4].height = 24
    
    act_group = {}
    total_general_horas = sum(float(r["horas"]) for r in registros) if registros else 0
    
    for reg in registros:
        act = reg["actividad"].strip() or "Sin Actividad"
        lugar = reg["lugar"].strip() or "Sin Especificar"
        if act not in act_group:
            act_group[act] = {
                "actividad": act,
                "lugares": set(),
                "fechas": set(),
                "horas": 0.0,
                "conteo": 0
            }
        act_group[act]["lugares"].add(lugar)
        act_group[act]["fechas"].add(reg["fecha"])
        act_group[act]["horas"] += float(reg["horas"])
        act_group[act]["conteo"] += 1
        
    act_list = sorted(act_group.values(), key=lambda x: x["horas"], reverse=True)
    
    row_idx2 = 5
    for item in act_list:
        ws2.cell(row=row_idx2, column=1, value=item["actividad"])
        
        lugares_str = ", ".join(sorted(list(item["lugares"])))
        ws2.cell(row=row_idx2, column=2, value=lugares_str)
        
        dias_str = f"{len(item['fechas'])} día(s) ({item['conteo']} reg.)"
        ws2.cell(row=row_idx2, column=3, value=dias_str).alignment = Alignment(horizontal="center")
        
        h_cell = ws2.cell(row=row_idx2, column=4, value=item["horas"])
        h_cell.number_format = '#,##0.00'
        h_cell.alignment = Alignment(horizontal="center")
        
        pct_val = (item["horas"] / total_general_horas) if total_general_horas > 0 else 0.0
        pct_cell = ws2.cell(row=row_idx2, column=5, value=pct_val)
        pct_cell.number_format = '0.0%'
        pct_cell.alignment = Alignment(horizontal="center")
        
        for col_num in range(1, 6):
            c = ws2.cell(row=row_idx2, column=col_num)
            c.font = data_font
            c.border = thin_border
            if row_idx2 % 2 == 0:
                c.fill = alt_fill
                
        ws2.row_dimensions[row_idx2].height = 20
        row_idx2 += 1
        
    if len(act_list) > 0:
        ws2.merge_cells(f"A{row_idx2}:C{row_idx2}")
        tot_label2 = ws2.cell(row=row_idx2, column=1, value="TOTAL HORAS POR ACTIVIDAD")
        tot_label2.font = bold_font
        tot_label2.alignment = Alignment(horizontal="right", vertical="center")
        
        tot_sum2 = ws2.cell(row=row_idx2, column=4, value=f"=SUM(D5:D{row_idx2-1})")
        tot_sum2.font = bold_font
        tot_sum2.number_format = '#,##0.00'
        tot_sum2.alignment = Alignment(horizontal="center", vertical="center")
        
        tot_pct2 = ws2.cell(row=row_idx2, column=5, value=1.0)
        tot_pct2.font = bold_font
        tot_pct2.number_format = '0.0%'
        tot_pct2.alignment = Alignment(horizontal="center", vertical="center")
        
        for col_num in range(1, 6):
            c = ws2.cell(row=row_idx2, column=col_num)
            c.fill = total_fill
            c.border = Border(top=Side(style='medium', color='64748B'), bottom=Side(style='double', color='1E293B'))
        ws2.row_dimensions[row_idx2].height = 25
        
    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1 and col_letter == 'A':
                continue
            try:
                if len(str(cell.value or '')) > max_len:
                    max_len = len(str(cell.value or ''))
            except:
                pass
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    ws2.column_dimensions['A'].width = 32
    ws2.column_dimensions['B'].width = 35
    ws2.column_dimensions['C'].width = 22
    ws2.column_dimensions['D'].width = 16
    ws2.column_dimensions['E'].width = 14
    
    # ---------------------------------------------------------
    # HOJA 3: RESUMEN POR CLIENTE / TAREA (SIN REPETIR CLIENTES/LUGAR)
    # ---------------------------------------------------------
    ws3 = wb.create_sheet(title="Horas por Cliente")
    
    title_fill3 = PatternFill(start_color="3B82F6", end_color="2563EB", fill_type="solid")
    ws3.merge_cells("A1:E1")
    title_cell3 = ws3["A1"]
    title_cell3.value = "RESUMEN AGRUPADO POR CLIENTE / TAREA (SIN REPETIR)"
    title_cell3.font = title_font
    title_cell3.fill = title_fill3
    title_cell3.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 35
    
    ws3["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Muestra cada cliente o tarea (ej. Interno) con todas sus actividades unidas en una sola fila"
    ws3["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    
    headers3 = ["Lugar / Cliente / Tarea", "Actividades Realizadas (Unidas)", "Días / Sesiones", "Horas Totales", "% del Total"]
    for col_num, header_text in enumerate(headers3, 1):
        cell = ws3.cell(row=4, column=col_num)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_num in [3, 4, 5] else "left", vertical="center")
        cell.border = thin_border
    ws3.row_dimensions[4].height = 24
    
    cli_pure_group = {}
    for reg in registros:
        lugar = reg["lugar"].strip() or "Sin Especificar"
        act = reg["actividad"].strip() or "Sin Actividad"
        if lugar not in cli_pure_group:
            cli_pure_group[lugar] = {
                "lugar": lugar,
                "actividades": set(),
                "fechas": set(),
                "horas": 0.0,
                "conteo": 0
            }
        cli_pure_group[lugar]["actividades"].add(act)
        cli_pure_group[lugar]["fechas"].add(reg["fecha"])
        cli_pure_group[lugar]["horas"] += float(reg["horas"])
        cli_pure_group[lugar]["conteo"] += 1
        
    cli_pure_list = sorted(cli_pure_group.values(), key=lambda x: x["horas"], reverse=True)
    
    row_idx3 = 5
    for item in cli_pure_list:
        ws3.cell(row=row_idx3, column=1, value=item["lugar"])
        
        acts_str = ", ".join(sorted(list(item["actividades"])))
        ws3.cell(row=row_idx3, column=2, value=acts_str)
        
        dias_str = f"{len(item['fechas'])} día(s) ({item['conteo']} reg.)"
        ws3.cell(row=row_idx3, column=3, value=dias_str).alignment = Alignment(horizontal="center")
        
        h_cell = ws3.cell(row=row_idx3, column=4, value=item["horas"])
        h_cell.number_format = '#,##0.00'
        h_cell.alignment = Alignment(horizontal="center")
        
        pct_val = (item["horas"] / total_general_horas) if total_general_horas > 0 else 0.0
        pct_cell = ws3.cell(row=row_idx3, column=5, value=pct_val)
        pct_cell.number_format = '0.0%'
        pct_cell.alignment = Alignment(horizontal="center")
        
        for col_num in range(1, 6):
            c = ws3.cell(row=row_idx3, column=col_num)
            c.font = data_font
            c.border = thin_border
            if row_idx3 % 2 == 0:
                c.fill = alt_fill
                
        ws3.row_dimensions[row_idx3].height = 20
        row_idx3 += 1
        
    if len(cli_pure_list) > 0:
        ws3.merge_cells(f"A{row_idx3}:C{row_idx3}")
        tot_label3 = ws3.cell(row=row_idx3, column=1, value="TOTAL HORAS POR CLIENTE / TAREA")
        tot_label3.font = bold_font
        tot_label3.alignment = Alignment(horizontal="right", vertical="center")
        
        tot_sum3 = ws3.cell(row=row_idx3, column=4, value=f"=SUM(D5:D{row_idx3-1})")
        tot_sum3.font = bold_font
        tot_sum3.number_format = '#,##0.00'
        tot_sum3.alignment = Alignment(horizontal="center", vertical="center")
        
        tot_pct3 = ws3.cell(row=row_idx3, column=5, value=1.0)
        tot_pct3.font = bold_font
        tot_pct3.number_format = '0.0%'
        tot_pct3.alignment = Alignment(horizontal="center", vertical="center")
        
        for col_num in range(1, 6):
            c = ws3.cell(row=row_idx3, column=col_num)
            c.fill = total_fill
            c.border = Border(top=Side(style='medium', color='64748B'), bottom=Side(style='double', color='1E293B'))
        ws3.row_dimensions[row_idx3].height = 25
        
    for col in ws3.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1 and col_letter == 'A':
                continue
            try:
                if len(str(cell.value or '')) > max_len:
                    max_len = len(str(cell.value or ''))
            except:
                pass
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    ws3.column_dimensions['A'].width = 32
    ws3.column_dimensions['B'].width = 38
    ws3.column_dimensions['C'].width = 22
    ws3.column_dimensions['D'].width = 16
    ws3.column_dimensions['E'].width = 14
    
    # ---------------------------------------------------------
    # HOJA 4: DESGLOSE POR CLIENTE / PROYECTO Y ACTIVIDAD (MATRIZ CRUZADA)
    # ---------------------------------------------------------
    ws4 = wb.create_sheet(title="Desglose Cliente x Act")
    
    title_fill4 = PatternFill(start_color="8B5CF6", end_color="6D28D9", fill_type="solid")
    ws4.merge_cells("A1:E1")
    title_cell4 = ws4["A1"]
    title_cell4.value = "DESGLOSE POR CLIENTE / PROYECTO Y ACTIVIDAD"
    title_cell4.font = title_font
    title_cell4.fill = title_fill4
    title_cell4.alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 35
    
    ws4["A2"] = f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Vista detallada por combinación Cliente + Actividad"
    ws4["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    
    headers4 = ["Lugar / Cliente / Proyecto", "Actividad", "Días / Sesiones", "Horas Totales", "% del Total"]
    for col_num, header_text in enumerate(headers4, 1):
        cell = ws4.cell(row=4, column=col_num)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_num in [3, 4, 5] else "left", vertical="center")
        cell.border = thin_border
    ws4.row_dimensions[4].height = 24
    
    cli_group = {}
    for reg in registros:
        lugar = reg["lugar"].strip() or "Sin Especificar"
        act = reg["actividad"].strip() or "Sin Actividad"
        key = (lugar, act)
        if key not in cli_group:
            cli_group[key] = {
                "lugar": lugar,
                "actividad": act,
                "fechas": set(),
                "horas": 0.0,
                "conteo": 0
            }
        cli_group[key]["fechas"].add(reg["fecha"])
        cli_group[key]["horas"] += float(reg["horas"])
        cli_group[key]["conteo"] += 1
        
    cli_list = sorted(cli_group.values(), key=lambda x: (x["lugar"], -x["horas"]))
    
    row_idx4 = 5
    for item in cli_list:
        ws4.cell(row=row_idx4, column=1, value=item["lugar"])
        ws4.cell(row=row_idx4, column=2, value=item["actividad"])
        
        dias_str = f"{len(item['fechas'])} día(s) ({item['conteo']} reg.)"
        ws4.cell(row=row_idx4, column=3, value=dias_str).alignment = Alignment(horizontal="center")
        
        h_cell = ws4.cell(row=row_idx4, column=4, value=item["horas"])
        h_cell.number_format = '#,##0.00'
        h_cell.alignment = Alignment(horizontal="center")
        
        pct_val = (item["horas"] / total_general_horas) if total_general_horas > 0 else 0.0
        pct_cell = ws4.cell(row=row_idx4, column=5, value=pct_val)
        pct_cell.number_format = '0.0%'
        pct_cell.alignment = Alignment(horizontal="center")
        
        for col_num in range(1, 6):
            c = ws4.cell(row=row_idx4, column=col_num)
            c.font = data_font
            c.border = thin_border
            if row_idx4 % 2 == 0:
                c.fill = alt_fill
                
        ws4.row_dimensions[row_idx4].height = 20
        row_idx4 += 1
        
    if len(cli_list) > 0:
        ws4.merge_cells(f"A{row_idx4}:C{row_idx4}")
        tot_label4 = ws4.cell(row=row_idx4, column=1, value="TOTAL HORAS POR CLIENTE / ACTIVIDAD")
        tot_label4.font = bold_font
        tot_label4.alignment = Alignment(horizontal="right", vertical="center")
        
        tot_sum4 = ws4.cell(row=row_idx4, column=4, value=f"=SUM(D5:D{row_idx4-1})")
        tot_sum4.font = bold_font
        tot_sum4.number_format = '#,##0.00'
        tot_sum4.alignment = Alignment(horizontal="center", vertical="center")
        
        tot_pct4 = ws4.cell(row=row_idx4, column=5, value=1.0)
        tot_pct4.font = bold_font
        tot_pct4.number_format = '0.0%'
        tot_pct4.alignment = Alignment(horizontal="center", vertical="center")
        
        for col_num in range(1, 6):
            c = ws4.cell(row=row_idx4, column=col_num)
            c.fill = total_fill
            c.border = Border(top=Side(style='medium', color='64748B'), bottom=Side(style='double', color='1E293B'))
        ws4.row_dimensions[row_idx4].height = 25
        
    for col in ws4.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1 and col_letter == 'A':
                continue
            try:
                if len(str(cell.value or '')) > max_len:
                    max_len = len(str(cell.value or ''))
            except:
                pass
        ws4.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    ws4.column_dimensions['A'].width = 32
    ws4.column_dimensions['B'].width = 32
    ws4.column_dimensions['C'].width = 22
    ws4.column_dimensions['D'].width = 16
    ws4.column_dimensions['E'].width = 14
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Reporte_Horas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

@app.route('/api/export/csv', methods=['GET'])
def export_csv():
    fecha_inicio = request.args.get('fecha_inicio')
    fecha_fin = request.args.get('fecha_fin')
    lugar = request.args.get('lugar')
    actividad = request.args.get('actividad')
    search = request.args.get('search')
    
    registros = db.get_all_registros(fecha_inicio, fecha_fin, lugar, actividad, search)
    feriados_list = db.get_feriados_chile()
    feriados_map = {f["fecha"]: f for f in feriados_list}
    
    for r in registros:
        f_info = feriados_map.get(r["fecha"])
        if f_info:
            prefix = f"[🇨🇱 Feriado: {f_info['titulo']}]"
            r["comentario"] = f"{prefix} {r['comentario'] or ''}".strip()
            
    df = pd.DataFrame(registros)
    if not df.empty:
        df = df[["id", "fecha", "lugar", "actividad", "horas", "comentario", "creado_en"]]
        df.columns = ["ID", "Fecha", "Lugar/Cliente", "Actividad", "Horas", "Comentario Complementario", "Fecha Registro"]
    else:
        df = pd.DataFrame(columns=["ID", "Fecha", "Lugar/Cliente", "Actividad", "Horas", "Comentario Complementario", "Fecha Registro"])
        
    output = io.BytesIO()
    df.to_csv(output, index=False, encoding='utf-8-sig')
    output.seek(0)
    
    filename = f"Reporte_Horas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return send_file(
        output,
        mimetype="text/csv",
        as_attachment=True,
        download_name=filename
    )

def abir_navegador():
    time.sleep(1.5)
    webbrowser.open("http://127.0.0.1:5000")

if __name__ == '__main__':
    try:
        import sys
        # Redirigir stdout/stderr a un archivo log si no hay consola (como en pythonw o exe sin consola)
        if sys.stdout is None or not sys.stdout.isatty():
            log_file = open(os.path.join(db.BASE_DIR, "servidor_fondo.log"), "a", encoding="utf-8", buffering=1)
            sys.stdout = log_file
            sys.stderr = log_file

        # Abrir navegador automáticamente al iniciar
        threading.Thread(target=abir_navegador, daemon=True).start()
        app.run(host='0.0.0.0', debug=False, port=5000)
    except BaseException as e:
        import traceback
        with open(os.path.join(db.BASE_DIR, "app_crash.log"), "w", encoding="utf-8") as f:
            f.write(traceback.format_exc())
